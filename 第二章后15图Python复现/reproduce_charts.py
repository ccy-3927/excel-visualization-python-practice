# -*- coding: utf-8 -*-
"""复现《第二章 图表(后15).xlsx》：15 张主图、2 张附加图，仅输出 PNG。

python reproduce_charts.py --input "完整路径/第二章 图表(后15).xlsx"
数据、图表关系、配色、图形填充和文本框均从源工作簿读取。
原工作簿只读，不执行工作簿中的文字说明，不生成 SVG。
"""
from pathlib import Path
from io import BytesIO
import argparse
import json
import math
import posixpath
import zipfile
import numpy as np
import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Wedge, Circle, Rectangle, PathPatch, Ellipse
from matplotlib.path import Path as MplPath
from matplotlib.lines import Line2D
from PIL import Image, ImageDraw, ImageFont
from chart_style import (xml, attr, value, fill, setup_font, native_size, text_shape,
                         draw_shapes, axes_for, vertical_axes, labels_above, BG, BLUE, PINK, WHITE)

HERE = Path(__file__).resolve().parent
FILENAME = '第二章 图表(后15).xlsx'
# 原文件把第六页命名为“20 南丁格尔（PPT）”，保留原名，避免误改编号。
MAIN_CHARTS = [1,2,3,4,5,6,7,8,10,12,13,14,15,16,17]


class Workbook:
    def __init__(self, path):
        self.path = path
        self.z = zipfile.ZipFile(path)
        self.book = openpyxl.load_workbook(path, data_only=True)

    def root(self, part):
        return xml(self.z.read(part))

    def rels(self, part):
        folder, name = posixpath.split(part)
        rel = folder + '/_rels/' + name + '.rels'
        if rel not in self.z.namelist():
            return {}
        return {e.get('Id'):posixpath.normpath(posixpath.join(folder,e.get('Target')))
                for e in self.root(rel)}

    def chart(self, cid):
        return self.root(f'xl/charts/chart{cid}.xml')

    def shapes(self, cid):
        part=f'xl/charts/chart{cid}.xml'
        rid=attr(self.root(part),'userShapes',key='id')
        return self.root(self.rels(part)[rid]) if rid else None

    def ref(self, formula):
        """通过图表真实引用取值，包括跨工作表引用和空白占位行。"""
        sheet, cells = formula.rsplit('!',1)
        sheet=sheet.strip("'").replace("''", "'")
        region=self.book[sheet][cells.replace('$','')]
        if isinstance(region,tuple):
            return [c.value for row in region for c in row]
        return [region.value]

    def series(self, ser, kind='val'):
        formula=ser.findtext(f'./{kind}/numRef/f') or ser.findtext(f'./{kind}/strRef/f')
        if formula:
            return self.ref(formula)
        return [float(e.text) for e in ser.findall(f'./{kind}/numLit/pt/v')]

    def col(self, sheet, column, first=3, last=8):
        return [self.book.worksheets[sheet-1].cell(r,column).value for r in range(first,last+1)]

    def picture(self,cid,rid):
        part=f'xl/charts/chart{cid}.xml'
        return np.asarray(Image.open(BytesIO(self.z.read(self.rels(part)[rid]))).convert('RGBA'))


def point_style(ser,index):
    for p in ser.findall('dPt'):
        if int(attr(p,'idx',-1))==index:
            sp=p.find('spPr')
            if sp is not None:
                return sp
    return ser.find('spPr')


def legend(fig,root,handles):
    node=root.find('./chart/legend')
    if node is None:
        return
    layout=node.find('./layout/manualLayout')
    fs=float(attr(node,'./txPr/p/pPr/defRPr',900,key='sz'))/100
    fig.legend(handles=handles,loc='upper left',
               bbox_to_anchor=(value(layout,'x',.1),1-value(layout,'y',.22)),
               ncol=len(handles),frameon=False,fontsize=fs,borderaxespad=0,
               handlelength=1.1,handletextpad=.3,columnspacing=.8)


def vertical(ax,root,cat,maximum,show_y=False):
    vertical_axes(ax,root,cat,maximum,200 if maximum<=1200 else 1000)
    if not show_y:
        ax.set_yticks([])


def second_axis(ax,lo,hi):
    right=ax.twinx();right.set_ylim(lo,hi);right.set_yticks([])
    right.patch.set_visible(False)
    for spine in right.spines.values():spine.set_visible(False)
    return right


def percentage_labels(ax,x,y,offset=5,size=9):
    for xx,yy in zip(x,y):
        ax.annotate(f'{yy:.0%}',(xx,yy),xytext=(0,offset),textcoords='offset points',
                    ha='center',va='bottom',fontsize=size,color=WHITE)


def circular_series(ax,root,book,cid):
    """按原图每个系列的内外环、颜色和透明占位复现多层扇形。"""
    chart=root.find('./chart/plotArea/doughnutChart')
    if chart is None:chart=root.find('./chart/plotArea/pieChart')
    series=chart.findall('ser')
    hole=value(chart,'holeSize',0)/100
    thickness=(1-hole)/len(series)
    start=90-value(chart,'firstSliceAng',0)
    endpoints=[]
    for j,ser in enumerate(series):
        vals=np.array(book.series(ser),float)
        angles=vals/vals.sum()*360
        radius=hole+(j+1)*thickness
        theta=start
        for k,sweep in enumerate(angles):
            style=point_style(ser,k);face=fill(style,BLUE)
            edge=fill(style.find('ln') if style is not None else None,'#00000000')
            lw=float(attr(style,'ln',0,key='w'))/12700
            if face[3] and sweep and cid not in (4,5):
                ax.add_patch(Wedge((0,0),radius,theta-sweep,theta,width=thickness,
                                   facecolor=face,edgecolor=edge,lw=lw))
            if k==0:endpoints.append((radius-thickness/2,theta-sweep,face))
            theta-=sweep
    ax.set_xlim(-1,1);ax.set_ylim(-1,1);ax.set_aspect('equal');ax.set_axis_off()
    if cid==2:
        # 玉玦图的角度以 50% 作满环，标签仍显示原始人数占比。
        for ser,(radius,theta,face) in zip(series,endpoints):
            pct=book.series(ser)[0]
            ang=np.deg2rad(theta+8)
            ax.text(radius*np.cos(ang),radius*np.sin(ang),f'{pct:.1%}',fontsize=8,
                    ha='center',va='center',color=WHITE)
    if cid in (4,5):
        ser=series[-1]
        vals=np.array(book.series(ser),float)
        cat=book.series(ser,'cat')
        theta=start
        for i,(name,v) in enumerate(zip(cat,vals)):
            angle=v/vals.sum()*360
            mid=np.deg2rad(theta-angle/2)
            # 标签与各自最外可见层对齐。
            visible=[hole+(j+1)*thickness for j,s in enumerate(series)
                     if fill(point_style(s,i),BLUE)[3]>0 and
                     not np.allclose(fill(point_style(s,i),BLUE)[:3],np.array([26,30,67])/255)]
            radius=max(visible) if visible else 1
            # 合并同色相邻环，避免 PNG 抗锯齿在无边框处产生细缝。
            sector_color=fill(point_style(series[0],i),BLUE)
            ax.add_patch(Wedge((0,0),radius,theta-angle,theta,width=radius-hole,
                               facecolor=sector_color,edgecolor='none'))
            px,py=radius*np.cos(mid),radius*np.sin(mid)
            tx=(radius+.18)*np.cos(mid);ty=(radius+.18)*np.sin(mid)
            if cid==4 and i==4:tx,ty=-.72,.46
            if cid==4 and i==5:tx,ty=-.18,.66
            if cid==5 and i==1:tx,ty=-.65,-.72
            ax.annotate(f'{str(name).strip()} {v:.1%}',(px,py),(tx,ty),
                        ha='left' if tx>=0 else 'right',va='center',fontsize=9,
                        arrowprops=dict(arrowstyle='-',lw=.5,color=WHITE),annotation_clip=False)
            theta-=angle


def draw_ppt_group(fig,group,transform=None):
    """保留 PPT 组合中的原始贝塞尔路径，递归应用组缩放和平移。"""
    xf=group.find('./grpSpPr/xfrm')
    off=xf.find('off');ext=xf.find('ext');co=xf.find('chOff');ce=xf.find('chExt')
    if transform is None:
        sx=1/float(ce.get('cx'));sy=1/float(ce.get('cy'))
        tx=-float(co.get('x'))*sx;ty=-float(co.get('y'))*sy
    else:
        psx,psy,ptx,pty=transform
        sx=psx*float(ext.get('cx'))/float(ce.get('cx'))
        sy=psy*float(ext.get('cy'))/float(ce.get('cy'))
        tx=ptx+psx*float(off.get('x'))-sx*float(co.get('x'))
        ty=pty+psy*float(off.get('y'))-sy*float(co.get('y'))
    for child in group:
        if child.tag=='grpSp':
            draw_ppt_group(fig,child,(sx,sy,tx,ty))
        elif child.tag=='sp':
            sp=child.find('spPr');xx=sp.find('xfrm')
            if xx is None:continue
            o,e=xx.find('off'),xx.find('ext')
            x,y=tx+sx*float(o.get('x')),ty+sy*float(o.get('y'))
            w,h=sx*float(e.get('cx')),sy*float(e.get('cy'))
            for path in sp.findall('./custGeom/pathLst/path'):
                pw,ph=float(path.get('w')),float(path.get('h'))
                verts,codes=[],[]
                for op in path:
                    if op.tag=='close':
                        verts.append(verts[0]);codes.append(MplPath.CLOSEPOLY);continue
                    code={'moveTo':MplPath.MOVETO,'lnTo':MplPath.LINETO,
                          'cubicBezTo':MplPath.CURVE4,'quadBezTo':MplPath.CURVE3}.get(op.tag)
                    if code is None:raise ValueError('未支持的 PPT 路径命令：'+op.tag)
                    for pt in op.findall('pt'):
                        verts.append((x+w*float(pt.get('x'))/pw,1-y-h*float(pt.get('y'))/ph))
                        codes.append(code)
                fig.add_artist(PathPatch(MplPath(verts,codes),transform=fig.transFigure,
                                         facecolor=fill(sp,BLUE),edgecolor='none',zorder=3))
            text_shape(fig,child,(x,y,w,h))


def render(book,sheet,cid,dpi,extra=False):
    root=book.chart(cid);shapes=book.shapes(cid)
    size=native_size(shapes) if cid not in (1,6,9,11) else (5.667,4.25)
    fig=plt.figure(figsize=size,dpi=dpi,facecolor=BG)
    ax=axes_for(fig,root)
    first=4 if sheet==10 or sheet==15 else 3
    last=7 if sheet==14 else 9 if sheet==10 else 8
    if sheet in (2,5):last=6
    cat=book.col(sheet,2,first,last)
    data=book.col(sheet,3,first,last)
    x=np.arange(len(cat))
    if cid==1:
        # 原图用波浪图片填充柱形，再用透明圆形遮罩裁切。
        ax.set_position([.285,.17,.45,.60]);ax.set_axis_off();ax.set_aspect('equal')
        ax.set_xlim(-1,1);ax.set_ylim(-1,1)
        rate=float(book.book.worksheets[0]['B4'].value)
        circle=Circle((0,0),.98,facecolor='none',edgecolor='none');ax.add_patch(circle)
        pic=book.picture(1,'rId1')
        water=ax.imshow(pic,extent=(-1,1,-1,-1+2*rate),origin='upper',interpolation='bicubic')
        water.set_clip_path(circle)
        ax.text(0,0,f'{rate:.0%}',fontsize=32,color=WHITE,ha='center',va='center')
        drawing=book.root('xl/drawings/drawing1.xml')
        for shape in drawing.findall('./twoCellAnchor/sp'):
            words=''.join(t.text or '' for t in shape.findall('.//t'))
            if '本科及以上学历' in words:text_shape(fig,shape,(.04,.04,.92,.23))
            elif '注：' in words:text_shape(fig,shape,(.035,.91,.9,.07))
    elif cid in (2,3,4,5):
        circular_series(ax,root,book,cid)
    elif cid==6:
        # PPT 页：颜色扇形是手工路径，透明 Excel 图表仅提供标签。
        ax.set_axis_off()
        group=book.root('xl/drawings/drawing11.xml').find('./twoCellAnchor/grpSp')
        draw_ppt_group(fig,group)
        drawing=book.root('xl/drawings/drawing11.xml')
        for shape in drawing.findall('./twoCellAnchor/sp'):
            words=''.join(t.text or '' for t in shape.findall('.//t'))
            if not words:continue
            xf=shape.find('./spPr/xfrm');off=xf.find('off');ext=xf.find('ext')
            box=((float(off.get('x'))-2590800)/5181600,
                 (float(off.get('y'))-176213)/3876675,
                 float(ext.get('cx'))/5181600,float(ext.get('cy'))/3876675)
            text_shape(fig,shape,box)
        ser=root.findall('.//ser')[-1]
        vals=book.series(ser);names=book.series(ser,'cat')
        # 按原图透明标签层引用第 19 页的数据，而非擅自改成本页数据。
        angles=90-np.cumsum([0]+[v/sum(vals)*360 for v in vals])
        for i,(name,v) in enumerate(zip(names,vals)):
            mid=np.deg2rad((angles[i]+angles[i+1])/2)
            fx=.55+.30*np.cos(mid);fy=.44+.31*np.sin(mid)
            fig.text(fx,fy,f'{str(name).strip()} {v:.1%}',fontsize=9,
                     ha='left' if np.cos(mid)>0 else 'right',va='center',zorder=15)
    elif cid==7:
        ax.set_axis_off();ax.set_aspect('equal');ax.set_xlim(-1,1);ax.set_ylim(-1,1)
        gauge=book.book.worksheets[6];score=float(gauge['H3'].value)
        # 三个环：内层标签、分段状态环、外层蓝色装饰环。
        ax.add_patch(Wedge((0,0),1,0,360,width=.1,facecolor=BLUE,edgecolor=BG,lw=.5))
        for k in range(10):
            end=225-27*k;start=end-27
            color='#087E8B' if k<3 else '#FFC54D' if k<7 else '#F94C66'
            ax.add_patch(Wedge((0,0),.9,start,end,width=.1,facecolor=color,edgecolor=BG,lw=.5))
        for k,scoretick in enumerate(range(50,151,10)):
            angle=np.deg2rad(225-27*k)
            ax.text(.745*np.cos(angle),.745*np.sin(angle),str(scoretick),ha='center',va='center',fontsize=8)
        angle=np.deg2rad(225-(score-50)/100*270)
        ax.plot([0,.84*np.cos(angle)],[0,.84*np.sin(angle)],color=WHITE,lw=1)
        # 中心数值和标题使用原图的文本框。
    elif cid==8:
        vertical(ax,root,cat,6000)
        ax.bar(x,data,width=1/3.19,color=BLUE)
        labels_above(ax,x,data)
        yoy=book.col(sheet,4)
        right=second_axis(ax,-1.2,.5)
        right.plot(x,yoy,color=PINK,lw=1.5,marker='o',ms=5)
        percentage_labels(right,x,yoy)
    elif cid in (9,11):
        # 工作表额外保留的制作过程图：保持其原始叠加与 1000 上限。
        vertical(ax,root,cat,1000,True)
        target=book.col(sheet,4,first,last)
        ax.bar(x,data,width=1/4.5,color=BLUE)
        ax.bar(x,target,bottom=data,width=1/4.5,color='none',edgecolor=BLUE)
    elif cid==10:
        vertical(ax,root,cat,1000)
        target=book.col(sheet,4)
        ax.bar(x,target,width=.25,color='none',edgecolor=(.95,.95,.95,.5),lw=.75)
        ax.bar(x,data,width=1/5.5,color=BLUE,alpha=.97)
        labels_above(ax,x,data)
        legend(fig,root,[Rectangle((0,0),1,1,facecolor='none',edgecolor=WHITE,label='目标销量'),
                         Rectangle((0,0),1,1,color=BLUE,label='实际销量')])
    elif cid==12:
        vertical(ax,root,cat,1200,True)
        base=np.zeros(len(cat));handles=[]
        for col,color,name in [(5,'#82ADD7','及格'),(6,(0,.69,.94,.7),'良好'),(7,BLUE,'优秀')]:
            vals=np.array(book.col(sheet,col,4,9))
            ax.bar(x,vals,bottom=base,width=1/3.5,color=color)
            base+=vals;handles.append(Rectangle((0,0),1,1,color=color,label=name))
        ax.bar(x,data,width=1/6,color='#0E5DFF')
        target=book.col(sheet,4,4,9)
        ax.plot(x,target,linestyle='none',marker='_',ms=10,mew=2,color='#FFC000')
        handles.extend([Rectangle((0,0),1,1,color='#0E5DFF',label='实际'),
                        Line2D([],[],linestyle='none',marker='_',color='#FFC000',label='目标')])
        legend(fig,root,handles)
    elif cid==13:
        vertical(ax,root,cat,5000)
        ax.bar(x,data,width=1/3.19,color=PINK);labels_above(ax,x,data)
        ypos=book.col(sheet,4);growth=book.col(sheet,5)
        points=root.findall('.//lineChart/ser/dPt')
        for i,(xx,yy,pct) in enumerate(zip(x,ypos,growth)):
            ms=value(points[i],'./marker/size',pct*100)
            ax.plot([xx],[yy],marker='o',ms=ms,color=BLUE,linestyle='none')
            ax.text(xx,yy,f'{pct:.0%}',ha='center',va='center',fontsize=8)
    elif cid==14:
        vertical(ax,root,cat,5000)
        previous=book.col(sheet,4);growth=book.col(sheet,5)
        width=1/(2+.27+2.19);delta=width*1.27/2
        ax.bar(x-delta,data,width,color=BLUE);ax.bar(x+delta,previous,width,color=PINK)
        labels_above(ax,x-delta,data);labels_above(ax,x+delta,previous)
        right=second_axis(ax,-1.1,.3)
        right.plot(x,growth,color='#FFC000',lw=1.5,marker='o',ms=5)
        percentage_labels(right,x,growth)
    elif cid==15:
        cat=book.col(sheet,2,3,17);data=book.col(sheet,3,3,17)
        cat=['' if c is None else c for c in cat];x=np.arange(15)
        vertical(ax,root,cat,12000)
        series=root.findall('.//barChart/ser');quarter=book.col(sheet,4,3,17)
        # 原图保留季度间隔行，季度总量是三根无间隙宽柱，月度柱覆盖在前。
        for i,v in enumerate(quarter):
            if v is not None:
                ax.bar(i,v,width=1,color=fill(point_style(series[1],i)),zorder=1)
        for i,v in enumerate(data):
            if v is not None:
                ax.bar(i,v,width=.4,color=fill(point_style(series[0],i)),zorder=3)
                labels_above(ax,[i],[v])
        for i in [1,5,9,13]:labels_above(ax,[i],[quarter[i]])
    elif cid in (16,17):
        ax.set_xlim(0,1);ax.set_ylim(-.5,len(cat)-.5);ax.set_xticks([])
        ax.set_yticks(x);ax.set_yticklabels(cat)
        if cid==16:
            bars=root.findall('.//barChart/ser')
            remaining=book.col(sheet,4,3,7)
            for i,(v,r) in enumerate(zip(data,remaining)):
                for ser,left,width in [(bars[0],0,v),(bars[1],v,r)]:
                    style=point_style(ser,i);rid=attr(style,'./blipFill/blip',key='embed')
                    ax.imshow(book.picture(cid,rid),extent=(left,left+width,i-.2,i+.2),
                              aspect='auto',interpolation='bicubic',zorder=2)
            ax.plot(data,x,ls='none',marker='o',ms=15,mfc=BLUE,mec=WHITE,mew=1,zorder=5)
            for i,v in enumerate(data):ax.text(v/2,i,f'{v:.0%}',ha='center',va='center',fontsize=9)
        else:
            old=book.col(sheet,4,4,8)
            ax.barh(x,np.ones(len(x)),height=1/6,color=(.95,.95,.95,.4),zorder=1)
            ax.barh(x,data,height=1/6,color=BLUE,zorder=2)
            ax.plot(old,x,ls='none',marker='o',ms=10,mfc='#A6A6A6',mec=WHITE,mew=1,zorder=4)
            ax.plot(data,x,ls='none',marker='o',ms=10,mfc=BLUE,mec=WHITE,mew=1,zorder=5)
            for i,(v,p) in enumerate(zip(data,old)):
                ax.annotate(f'{v:.0%}',(v,i),xytext=(0,9),textcoords='offset points',ha='center',fontsize=8)
                ax.annotate(f'{p:.0%}',(p,i),xytext=(0,-15),textcoords='offset points',ha='center',fontsize=8)
    if cid not in (1,6):
        draw_shapes(fig,shapes)
        # 原通用文本渲染器不处理椭圆；补充比较滑珠图的图例圆点。
        if shapes is not None:
            for a in shapes:
                if attr(a,'./sp/spPr/prstGeom',key='prst')=='ellipse':
                    x0,y0=float(a.findtext('./from/x')),float(a.findtext('./from/y'))
                    ww=float(a.findtext('./to/x'))-x0;hh=float(a.findtext('./to/y'))-y0
                    sp=a.find('./sp/spPr')
                    fig.add_artist(Ellipse((x0+ww/2,1-y0-hh/2),ww,hh,transform=fig.transFigure,
                                           facecolor=fill(sp),edgecolor=fill(sp.find('ln')),lw=.75,zorder=8))
    return fig


def overview(files,target):
    canvas=Image.new('RGB',(1800,2450),'#10142D');draw=ImageDraw.Draw(canvas)
    font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',18) if Path('C:/Windows/Fonts/msyh.ttc').exists() else ImageFont.load_default()
    for i,p in enumerate(files):
        im=Image.open(p).convert('RGB')
        im.thumbnail((580,445),Image.Resampling.LANCZOS if hasattr(Image,'Resampling') else Image.LANCZOS)
        x=(i%3)*600+(600-im.width)//2;y=(i//3)*490+30
        canvas.paste(im,(x,y));draw.text(((i%3)*600+12,(i//3)*490+5),p.stem,font=font,fill='white')
    canvas.save(target)


def validate(book):
    """核对图表缓存与实际单元格，记录源文件自身的差异。"""
    checked=0;differences=[]
    for cid in range(1,18):
        for ref in book.chart(cid).findall('.//numRef'):
            f=ref.findtext('f');cache=ref.find('numCache')
            if f is None or cache is None:continue
            values=book.ref(f)
            for pt in cache.findall('pt'):
                i=int(pt.get('idx'))
                if i>=len(values) or values[i] is None:continue
                a,b=float(values[i]),float(pt.findtext('v'))
                checked+=1
                if not math.isclose(a,b,rel_tol=1e-9,abs_tol=1e-9):
                    differences.append({'chart':cid,'reference':f,'index':i,'cell':a,'cache':b})
    return checked,differences


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input',type=Path)
    parser.add_argument('--output',type=Path,default=HERE/'生成图表')
    parser.add_argument('--dpi',type=int,default=240)
    args=parser.parse_args()
    candidates=[HERE/FILENAME,HERE.parent/FILENAME,Path.cwd()/FILENAME,
                Path(r'D:/CP/python可视化课程/《Excel数据可视化 - 从图表到数据大屏》-清华-郭宏远')/FILENAME]
    source=args.input or next((p for p in candidates if p.exists()),None)
    if source is None or not source.exists():parser.error('找不到工作簿，请使用 --input 指定路径。')
    if args.dpi<=0:parser.error('--dpi 必须为正数。')
    setup_font();book=Workbook(source)
    args.output.mkdir(parents=True,exist_ok=True)
    files=[]
    for sheet,cid in enumerate(MAIN_CHARTS,1):
        name=book.book.worksheets[sheet-1].title.strip().replace(' ','_',1)
        fig=render(book,sheet,cid,args.dpi)
        p=args.output/(name+'.png');fig.savefig(p,dpi=args.dpi,facecolor=BG);plt.close(fig)
        files.append(p);print(f'已生成 {sheet:02d}/15：{p.name}')
    extras=args.output/'附加图表';extras.mkdir(exist_ok=True)
    for sheet,cid,name in [(9,9,'24_目标柱形图_原始制作图'),(10,11,'25_子弹图_原始制作图')]:
        fig=render(book,sheet,cid,args.dpi,extra=True)
        fig.savefig(extras/(name+'.png'),dpi=args.dpi,facecolor=BG);plt.close(fig)
    overview(files,args.output/'00_后15图总览.png')
    checked,diffs=validate(book)
    report={'source':str(source.resolve()),'main_charts':15,'extra_charts':2,
            'image_format':'PNG','validated_values':checked,'source_cache_differences':diffs,
            'main_files':[p.name for p in files],
            'note':'Python 重绘，非 Excel 原生导出。PPT 页保留原始组合路径及跨工作表标签引用。'}
    (args.output/'生成记录.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    book.book.close();book.z.close()
    print(f'完成，核对 {checked} 处数值，缓存差异 {len(diffs)} 处。')
    print('输出目录：'+str(args.output.resolve()))


if __name__=='__main__':main()
