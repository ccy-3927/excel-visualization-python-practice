"""直接从 Excel 明细计算；不依赖旧 HTML 或旧版 reproduce.py。"""
from pathlib import Path
from collections import Counter, defaultdict
import openpyxl

HR_FILE = '第四章 人力资源可视化看板.xlsx'
SALES_FILE = '第四章 销售看板参考.xlsx'
DEPARTMENTS = ['销售部', '研发部', '信息技术部', '行政部', '人力资源部', '财务部']
REGIONS = ['华北', '华南', '东北', '西北', '西南', '华东']
COURIERS = ['顺丰', '韵达', '中通', '申通', '圆通', 'EMS']
CATEGORIES = ['办公用品', '家具产品', '数码电子']
HR_SPECS = [
    ('年龄', 7, ['18-24', '25-29', '30-34', '35-39', '40=<'], 'pie'),
    ('性别', 2, ['男', '女'], 'pie'),
    ('婚姻状况', 4, ['单身', '已婚', '离异'], 'pie'),
    ('学历', 3, ['专科以下', '专科', '本科', '硕士研究生', '博士研究生'], 'horizontal'),
    ('入转调', 6, ['入职', '转入', '转出'], 'bar'),
    ('部门', 5, DEPARTMENTS, 'horizontal'),
]

def read_data(directory):
    """读取真实源文件。缓存由 app 根据文件修改时间控制。"""
    directory = Path(directory)
    with_hr = openpyxl.load_workbook(directory / HR_FILE, data_only=True, read_only=True)
    try:
        hr = [tuple(r) for r in with_hr.worksheets[0].iter_rows(min_row=2, max_col=8, values_only=True) if r[0] is not None]
    finally:
        with_hr.close()
    with_sales = openpyxl.load_workbook(directory / SALES_FILE, data_only=True, read_only=True)
    try:
        sales = [tuple(r) for r in with_sales['销售明细'].iter_rows(min_row=2, max_col=10, values_only=True) if r[0] is not None]
        costs = [tuple(r) for r in with_sales['成本明细'].iter_rows(min_row=2, max_col=3, values_only=True) if r[0] is not None]
    finally:
        with_sales.close()
    return hr, sales, costs

def summarize_hr(rows, departments=None, genders=None):
    selected = [r for r in rows if (departments is None or r[5] in departments) and (genders is None or r[2] in genders)]
    groups = {}
    for title, column, labels, kind in HR_SPECS:
        counter = Counter(r[column] for r in selected)
        groups[title] = {label: counter[label] for label in labels}
    return {'count': len(selected), 'groups': groups}

def summarize_sales(rows, costs):
    monthly = {}
    for month in range(1, 13):
        selected = [r for r in rows if r[9] == f'{month}月']
        products = defaultdict(lambda: [0.0, 0])
        for row in selected:
            products[row[8]][0] += row[4]
            products[row[8]][1] += 1
        revenue = sum(r[4] for r in selected) / 10000
        profit = sum(r[6] for r in selected) / 10000
        cost = defaultdict(float)
        for m, category, value in costs:
            if m == f'{month}月':
                cost[category] += value
        monthly[month] = {
            'count': len(selected), 'revenue': revenue, 'profit': profit,
            'margin': profit / revenue if revenue else None,
            'region': dict(Counter(r[2] for r in selected)),
            'courier': dict(Counter(r[3] for r in selected)),
            'category': dict(Counter(r[7] for r in selected)), 'cost': dict(cost),
            'top': [{'产品': name, '销售额(万)': amount / 10000, '销量': count}
                    for name, (amount, count) in sorted(products.items(), key=lambda kv: (-kv[1][0], kv[0]))[:3]],
        }
    return monthly

def month_change(monthly, month):
    previous = monthly.get(month - 1)
    if not previous or not previous['count']:
        return None
    return monthly[month]['count'] / previous['count'] - 1

def next_month(month):
    return month % 12 + 1
