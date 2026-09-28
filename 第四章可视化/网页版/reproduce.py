"""第四章看板复现。Python 读取 Excel、聚合数据、绘制 SVG 并生成离线交互 HTML。
运行：python reproduce.py --source-dir 原Excel所在目录
依赖：openpyxl。浏览器交互仅使用内置 JavaScript，无网络/CDN/服务端依赖。
"""
from pathlib import Path
from collections import Counter, defaultdict
from html import escape
import argparse, json, math
import openpyxl

HR_FILE = '第四章 人力资源可视化看板.xlsx'
SALES_FILE = '第四章 销售看板参考.xlsx'
CYAN, PINK, BLUE, ORANGE = '#7BBDD5', '#FC6C9D', '#4A5BD1', '#F6A576'

def text(x,y,value,size=14,anchor='middle',color='#edf2ff'):
    return f'<text x="{x}" y="{y}" font-size="{size}" text-anchor="{anchor}" fill="{color}">{escape(str(value))}</text>'

def svg(body,w=400,h=200):
    return f'<svg viewBox="0 0 {w} {h}" role="img" xmlns="http://www.w3.org/2000/svg">{body}</svg>'

def donut(labels,values,colors,center='',subtitle='',compact=False):
    """SVG 圆环：每段含原生悬浮提示，图例同时给人数/占比。"""
    total=sum(values); cx,cy,r=(90,90,62) if compact else (108,112,78)
    body=f'<circle cx="{cx}" cy="{cy}" r="{r}" fill="none" stroke="#465477" stroke-width="17"/>'
    offset=0; circumference=2*math.pi*r
    for i,(label,value) in enumerate(zip(labels,values)):
        frac=value/total if total else 0
        if frac>0:
            a=offset*2*math.pi-math.pi/2; b=(offset+min(frac,.9999999))*2*math.pi-math.pi/2
            x1,y1=cx+r*math.cos(a),cy+r*math.sin(a);x2,y2=cx+r*math.cos(b),cy+r*math.sin(b)
            body+=f'<path class="arc" d="M{x1} {y1} A{r} {r} 0 {int(frac>.5)} 1 {x2} {y2}" fill="none" stroke="{colors[i%len(colors)]}" stroke-width="17"><title>{escape(str(label))}：{value:,.2f} / {frac:.2%}</title></path>'
        offset+=frac
        if not compact:
            yy=48+i*34
            body+=f'<rect x="215" y="{yy-10}" width="9" height="9" fill="{colors[i%len(colors)]}"/>'+text(232,yy,f'{label}  {value:,}  ({frac:.0%})',12,'start')
    body+=text(cx,cy+2,center,27)+text(cx,cy+27,subtitle,12)
    return svg(body,180,180) if compact else svg(body,400,225)

def bars(labels,values,color=CYAN,horizontal=False,colors=None):
    n=len(labels); body=''; mx=max(values+[1])*1.22
    if horizontal:
        for i,(label,v) in enumerate(zip(labels,values)):
            y=17+i*172/max(n,1); width=220*v/mx
            body+=text(112,y+14,label,12,'end')
            body+=f'<rect class="bar" x="122" y="{y}" width="{width}" height="19" fill="{(colors or [color]*n)[i]}"><title>{escape(str(label))}：{v:,}</title></rect>'+text(128+width,y+14,f'{v:,}',12,'start')
    else:
        body+='<path d="M22 165 H390" stroke="#5b6589" fill="none"/>'
        for i,(label,v) in enumerate(zip(labels,values)):
            x=30+i*355/n; height=v/mx*130; width=min(40,240/n)
            body+=f'<rect class="bar" x="{x}" y="{165-height}" width="{width}" height="{height}" fill="{color}"><title>{escape(str(label))}：{v:,}</title></rect>'+text(x+width/2,155-height,f'{v:,}',12)+text(x+width/2,188,label,12)
        body+=text(14,169,0,10)
    return svg(body)

def area(values,month):
    mx=math.ceil(max(values)/50)*50; body=''
    for tick in range(5):
        y=165-tick*34; body+=text(38,y+4,f'{mx*tick/4:.0f}',10,'end')+f'<path d="M45 {y} H388" stroke="#424d78"/>'
    pts=[(48+i*30.5,165-v/mx*136) for i,v in enumerate(values)]
    body+=f'<path d="M48 165 '+ ' '.join(f'L{x} {y}' for x,y in pts)+' L383.5 165Z" fill="#4A5BD1" opacity=".8"/>'
    for i,(x,y) in enumerate(pts):
        body+=f'<circle cx="{x}" cy="{y}" r="{5 if i+1==month else 3}" fill="{CYAN if i+1==month else BLUE}"><title>{i+1}月：{values[i]:.4f} 万</title></circle>'+text(x,187,f'{i+1}月',10)
    return svg(body)

def panel(title,body,extra=''):
    return f'<section class="panel {extra}"><h2>{title}</h2>{body}</section>'

CSS='''*{box-sizing:border-box}body{margin:0;background:#1b2047;color:#edf2ff;font-family:"Microsoft YaHei","Noto Sans CJK SC",sans-serif}main{max-width:1500px;margin:auto;padding:18px}header{display:flex;align-items:center;justify-content:space-between;background:#2c3560;padding:18px 24px;gap:20px}h1{font-size:25px;margin:0;color:#c6e9ce}h2{font-size:16px;text-align:center;font-weight:500;margin:0 0 12px;color:#e0edfd}nav{display:flex;gap:10px;align-items:center;flex-wrap:wrap}button,select,a{font:inherit}button,select{border:1px solid #7bbdd5;border-radius:5px;padding:7px 12px;color:#edf2ff;background:#232b52;cursor:pointer}a{color:#a6d8ef}nav a{font-size:13px}.sales{display:grid;grid-template-columns:1fr 1fr 1fr;gap:10px;margin-top:10px}.column{display:flex;flex-direction:column;gap:10px}.panel{background:#2c3560;padding:15px 12px;min-width:0}.panel svg{width:100%;display:block}.hr{display:grid;grid-template-columns:repeat(3,1fr);gap:12px;margin-top:14px}.hr .panel{background:#1d244a;min-height:300px;padding:20px}.costs{display:grid;grid-template-columns:1fr 1fr;gap:10px}.cost{border:1px solid #7bbdd5;border-radius:12px;padding:12px;min-width:0}.cost h3{margin:0;font-size:14px;font-weight:400}.cost .value{text-align:center;font-size:19px;color:#c6e9ce}.cost svg{width:100%;max-height:122px}.profit{display:grid;grid-template-columns:1fr 1fr;align-items:center}.profit p{font-size:13px}.profit strong{font-size:23px;color:#fc6c9d;display:block;margin:5px 0 16px}.change{font-size:23px;color:#c6e9ce;text-align:right;padding:0 10px}.change small{font-size:12px;color:#b6c3e5}.top{width:100%;border-collapse:collapse;font-size:12px}.top td,.top th{padding:14px 5px;text-align:left;border-bottom:1px solid #404b76}.top td:nth-child(3){color:#a8bbff}.top td:nth-child(2){max-width:170px;overflow-wrap:anywhere}.top th{font-weight:400;color:#b6c3e5}.rank{font-size:20px;color:#ffd077}footer{font-size:12px;line-height:1.9;color:#b6c3e5;background:#2c3560;padding:12px 20px;margin-top:10px}.kpi{color:#c6e9ce;font-size:23px}.bar,.arc{animation:reveal .55s ease-out}@keyframes reveal{from{opacity:.25}to{opacity:1}}@media(max-width:900px){.sales,.hr{grid-template-columns:1fr}header{flex-direction:column;align-items:flex-start}.hr .panel{min-height:0}}@media(prefers-reduced-motion:reduce){*{animation:none!important}}'''

def document(title,header,body,script=''):
    return f'<!DOCTYPE html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{title}</title><style>{CSS}</style><main>{header}{body}</main>{script}</html>'

def load_hr(path):
    w=openpyxl.load_workbook(path,data_only=True); s=w.worksheets[0]
    rows=[r for r in s.iter_rows(min_row=2,max_col=8,values_only=True) if r[0] is not None]
    specs=[('年龄',7,['18-24','25-29','30-34','35-39','40=<'],['#00B0F0','#ED7D31','#7AD5FE','#EDA530','#0070C0'],'ring'),('性别',2,['男','女'],['#EDA530','#0070C0'],'ring'),('婚姻状况',4,['单身','已婚','离异'],['#EDA530','#0070C0','#7AD5FE'],'ring'),('学历',3,['专科以下','专科','本科','硕士研究生','博士研究生'],[],'horizontal'),('入转调',6,['入职','转入','转出'],[],'bar'),('部门',5,['销售部','研发部','信息技术部','行政部','人力资源部','财务部'],[],'horizontal')]
    body='';counts={}
    for title,col,labels,colors,kind in specs:
        c=Counter(r[col] for r in rows); vals=[c[x] for x in labels];counts[title]=dict(zip(labels,vals))
        body+=panel(title,donut([x.replace('40=<','40岁及以上') for x in labels],vals,colors,str(sum(vals)),'人') if kind=='ring' else bars(labels,vals,'#7AD5FE',kind=='horizontal'))
    header=f'<header><h1>公司人员结构看板</h1><nav><span class="kpi">总人数：{len(rows):,}</span><a href="销售看板.html">销售看板 →</a></nav></header>'
    html=document('公司人员结构看板',header,f'<div class="hr">{body}</div><footer>数据：2022年3月人员基础信息。按原工作簿保留六张图的图型、顺序与配色。悬停图形可查看数值。<br>入转调仅统计本月有记录的人员；空白不算入职、转入或转出。原年龄标签“40=&lt;”显示为“40岁及以上”。</footer>')
    html=html.replace('</footer>','<br>原表学历公式 K24/K25 引用颠倒；此处按明细修正为：专科以下170人、专科282人。</footer>')
    return html,counts,len(rows)

def load_sales(path):
    w=openpyxl.load_workbook(path,data_only=True)
    rows=[r for r in w['销售明细'].iter_rows(min_row=2,max_col=10,values_only=True) if r[0] is not None]
    costs=[r for r in w['成本明细'].iter_rows(min_row=2,max_col=3,values_only=True) if r[0] is not None]
    monthly={}
    for m in range(1,13):
        rr=[r for r in rows if r[9]==f'{m}月']
        products=defaultdict(lambda:[0,0])
        for r in rr: products[r[8]][0]+=r[4]; products[r[8]][1]+=1
        revenue=sum(r[4] for r in rr)/10000; profit=sum(r[6] for r in rr)/10000
        monthly[m]={'month':m,'count':len(rr),'revenue':revenue,'profit':profit,'margin':profit/revenue if revenue else None,'region':dict(Counter(r[2] for r in rr)),'courier':dict(Counter(r[3] for r in rr)),'category':dict(Counter(r[7] for r in rr)),'cost':{r[1]:r[2] for r in costs if r[0]==f'{m}月'},'top':[(k,v[0]/10000,v[1]) for k,v in sorted(products.items(),key=lambda kv:(-kv[1][0],kv[0]))[:3]]}
    return monthly,w['统计数据'],len(rows)

def sales_frame(d,allmonths):
    m=d['month'];labels=['办公用品','家具产品','数码电子']
    left=panel('产品类别销量',donut(labels,[d['category'].get(x,0) for x in labels],[PINK,BLUE,ORANGE],str(d['count']),'总销量'))
    cs=''; total=sum(d['cost'].values())
    for label in ['推广','人工','产品','其他']:
        v=d['cost'][label];cs+=f'<div class="cost"><h3>{label}费用合计</h3>'+donut([label,'其他成本'],[v,total-v],[CYAN,'#465477'],f'{v/total:.2%}',compact=True)+f'<div class="value">{v:,.2f}</div></div>'
    left+=panel('成本费用',f'<div class="costs">{cs}</div>')
    prev=allmonths.get(m-1); rate=(d['count']/prev['count']-1) if prev and prev['count'] else None
    cmp=bars([f'{m}月',f'{m-1}月'] if prev else ['1月'],[d['count'],prev['count']] if prev else [d['count']],horizontal=True,colors=[CYAN,BLUE])
    middle=panel('销量环比',f'<div class="change">{format(rate,"+.2%") if rate is not None else "—"} <small>{"较上月" if prev else "无上月数据"}</small></div>'+cmp)
    table='<table class="top"><thead><tr><th>排名</th><th>产品</th><th>销售额(万)</th><th>销量</th></tr></thead><tbody>'
    for i,(name,v,n) in enumerate(d['top'],1):table+=f'<tr><td class="rank">{i}</td><td>{escape(name)}</td><td>{v:.2f}</td><td>{n}</td></tr>'
    middle+=panel('商品销售额 Top3',table+'</tbody></table>')
    margin=d['margin'];ring=donut(['利润','销售额减利润'],[d['profit'],d['revenue']-d['profit']],[PINK,'#465477'],f'{margin:.2%}','利润率',True) if margin is not None and 0<=margin<=1 else f'<div class="kpi">{format(margin,".2%") if margin is not None else "—"}</div>'
    middle+=panel('利润额占比销售额',f'<div class="profit">{ring}<div><p>利润额(万)<strong>{d["profit"]:.2f}</strong></p><p>销售额(万)<strong>{d["revenue"]:.2f}</strong></p></div></div>')
    right=panel('各月销售额(万)',area([allmonths[x]['revenue'] for x in range(1,13)],m))
    for title,key,labels,col in [('区域销量','region',['华北','华南','东北','西北','西南','华东'],CYAN),('快递公司销量','courier',['顺丰','韵达','中通','申通','圆通','EMS'],PINK)]:right+=panel(title,bars(labels,[d[key].get(x,0) for x in labels],col))
    return ''.join(f'<div class="column">{x}</div>' for x in [left,middle,right])

def sales_document(monthly):
    frames={m:sales_frame(d,monthly) for m,d in monthly.items()}
    controls='<nav><a href="人力资源看板.html">← 人力资源看板</a><label for="month">月份 </label><select id="month">'+''.join(f'<option value="{m}" {"selected" if m==9 else ""}>{m}月</option>' for m in range(1,13))+'</select><button id="play" aria-pressed="false">▶ 自动播放</button><button id="reset">恢复9月</button></nav>'
    header='<header><h1>公司销售数据可视化看板</h1>'+controls+'</header>'
    body=f'<div class="sales" id="dashboard" aria-live="polite">{frames[9]}</div><footer><span id="period">当前：2021年9月</span> · 销量按明细记录数统计，与原表 COUNTIF 口径一致。销售额与利润额单位为万；成本保留源表单位（未注明）。<br>原表“同比”实际计算本月与上月的环比，此处已更正；1月无上年12月数据，环比显示“—”。年度销售趋势保留全年数据，月份切换时突出当前月。自动播放为新增演示功能。</footer>'
    script='<script>const frames='+json.dumps(frames,ensure_ascii=False).replace('</',r'<\/')+''';
const sel=document.getElementById('month'), play=document.getElementById('play');let timer=null;
function render(){document.getElementById('dashboard').innerHTML=frames[sel.value];document.getElementById('period').textContent='当前：2021年'+sel.value+'月';}
function stop(){clearInterval(timer);timer=null;play.textContent='▶ 自动播放';play.setAttribute('aria-pressed','false');}
sel.addEventListener('change',render);
play.onclick=()=>{if(timer){stop();return;}play.textContent='❚❚ 暂停';play.setAttribute('aria-pressed','true');timer=setInterval(()=>{sel.value=Number(sel.value)%12+1;render();},1800);};
document.getElementById('reset').onclick=()=>{stop();sel.value=9;render();};
document.addEventListener('visibilitychange',()=>{if(document.hidden)stop();});
</script>'''
    return document('公司销售数据可视化看板',header,body,script)

def verify(monthly,s,hr,headcount,rowcount):
    """独立核对原 Excel 已保存的9月计算结果，以及全年各维度加总。"""
    checks=[]
    def check(name,a,b):
        ok=math.isclose(a,b,rel_tol=1e-9,abs_tol=1e-7);checks.append({'项目':name,'复现值':a,'Excel或明细值':b,'通过':ok})
        if not ok:raise ValueError(f'核对失败：{name}: {a} != {b}')
    d=monthly[9]
    for key,cell in [('count','I2'),('revenue','K2'),('profit','M2'),('margin','N3')]:check('9月 '+key,d[key],s[cell].value)
    check('9月成本',sum(d['cost'].values()),s['G2'].value)
    for key,labels,col,start in [('region',['华北','华南','东北','西北','西南','华东'],'E',9),('courier',['顺丰','韵达','中通','申通','圆通','EMS'],'E',18),('category',['办公用品','家具产品','数码电子'],'K',13)]:
        for i,label in enumerate(labels):check('9月 '+label,d[key].get(label,0),s[f'{col}{start+i}'].value)
    for i,(_,amount,count) in enumerate(d['top'],3):
        check(f'Top{i-2}金额',amount,s[f'Y{i}'].value);check(f'Top{i-2}销量',count,s[f'Z{i}'].value)
        assert d['top'][i-3][0]==s[f'X{i}'].value
    for m,v in monthly.items():
        check(f'{m}月销售额',v['revenue'],s[f'K{17+m}'].value)
        for key in ['region','courier','category']:check(f'{m}月{key}合计',sum(v[key].values()),v['count'])
    check('全年销量',sum(v['count'] for v in monthly.values()),rowcount)
    for key,c in hr.items():
        if key!='入转调':check('人员 '+key,sum(c.values()),headcount)
    check('总人数',headcount,1470)
    return checks

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-dir',type=Path,default=Path(__file__).resolve().parent.parent)
    parser.add_argument('--output-dir',type=Path,default=Path(__file__).resolve().parent)
    args=parser.parse_args();args.output_dir.mkdir(parents=True,exist_ok=True)
    for name in [HR_FILE,SALES_FILE]:
        if not (args.source_dir/name).is_file():parser.error(f'找不到 {args.source_dir/name}；请使用 --source-dir 指定Excel目录')
    hr,counts,n=load_hr(args.source_dir/HR_FILE);monthly,s,nrows=load_sales(args.source_dir/SALES_FILE)
    checks=verify(monthly,s,counts,n,nrows)
    for name,content in [('人力资源看板.html',hr),('销售看板.html',sales_document(monthly))]:
        (args.output_dir/name).write_text(content,encoding='utf-8')
    (args.output_dir/'核对报告.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
    (args.output_dir/'汇总数据.json').write_text(json.dumps({'人力资源':counts,'销售':monthly},ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'已生成两个看板；{len(checks)}项数值核对通过。输出：{args.output_dir}')

if __name__=='__main__':main()
