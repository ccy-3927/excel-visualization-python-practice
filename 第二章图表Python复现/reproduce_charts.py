# -*- coding: utf-8 -*-
"""读取《第二章 图表(前15).xlsx》，用 Matplotlib 复现 15 张示例图。

运行：python reproduce_charts.py
指定源文件：python reproduce_charts.py --input "D:/.../第二章 图表(前15).xlsx"
依赖：numpy、matplotlib、openpyxl、Pillow。无需启动 Excel。
样式、标题、绘图区位置与图片填充从 xlsx 内的 DrawingML 读取。
数据从单元格读取；不会执行工作簿中的说明文字或修改原文件。
"""
from pathlib import Path
from io import BytesIO
import argparse
import json
import math
import posixpath
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta

import numpy as np
import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager, colors
from matplotlib.patches import Rectangle, FancyBboxPatch, Wedge, Circle
from matplotlib.lines import Line2D
from PIL import Image, ImageDraw, ImageFont

BG = '#1A1E43'
BLUE, PINK, WHITE = '#0070C0', '#E74E69', '#F2F2F2'
FILENAME = '第二章 图表(前15).xlsx'
HERE = Path(__file__).resolve().parent


def xml(data):
    """去掉命名空间前缀，保留 XML 内容和关系标识。"""
    root = ET.fromstring(data)
    for e in root.iter():
        e.tag = e.tag.split('}')[-1]
        e.attrib = {k.split('}')[-1]: v for k, v in e.attrib.items()}
    return root


def attr(root, path, default=None, key='val'):
    e = root.find(path) if root is not None else None
    return e.get(key, default) if e is not None else default


def value(root, path, default=0):
    return float(attr(root, path, default))


def fill(root, default=WHITE):
    """读取 RGB、主题色、亮度及透明度。"""
    if root is None:
        return colors.to_rgba(default)
    if root.find('noFill') is not None:
        return (0, 0, 0, 0)
    solid = root.find('solidFill')
    if solid is None:
        return colors.to_rgba(default)
    e = next(iter(solid), None)
    if e is None:
        return colors.to_rgba(default)
    palette = {'bg1': '#FFFFFF', 'lt1': '#FFFFFF', 'tx1': '#000000',
               'dk1': '#000000', 'accent1': '#4472C4', 'accent2': '#ED7D31'}
    c = '#' + e.get('val') if e.tag == 'srgbClr' else palette.get(e.get('val'), default)
    rgb = np.array(colors.to_rgb(c))
    rgb = rgb * value(e, 'lumMod', 100000) / 100000 + value(e, 'lumOff', 0) / 100000
    return (*np.clip(rgb, 0, 1), value(e, 'alpha', 100000) / 100000)


class Source:
    def __init__(self, path):
        self.path = path
        self.z = zipfile.ZipFile(path)
        self.book = openpyxl.load_workbook(path, data_only=True)

    def root(self, part):
        return xml(self.z.read(part))

    def rels(self, part):
        folder, name = posixpath.split(part)
        rel = folder + '/_rels/' + name + '.rels'
        if rel not in self.z.namelist():
            return {}
        return {e.get('Id'): posixpath.normpath(posixpath.join(folder, e.get('Target')))
                for e in self.root(rel)}

    def col(self, sheet, column, first=3, last=None):
        s = self.book.worksheets[sheet - 1]
        if last is None:
            last = first
            while s.cell(last + 1, 2).value is not None:
                last += 1
        return [s.cell(r, column).value for r in range(first, last + 1)]

    def chart(self, sheet):
        return self.root('xl/charts/chart%d.xml' % (sheet if sheet < 7 else sheet - 1))

    def shapes(self, sheet):
        part = 'xl/charts/chart%d.xml' % (sheet if sheet < 7 else sheet - 1)
        r = self.root(part)
        rid = attr(r, 'userShapes', key='id')
        return self.root(self.rels(part)[rid]) if rid else None

    def picture(self, chart, rid):
        part = 'xl/charts/chart%d.xml' % chart
        return np.asarray(Image.open(BytesIO(self.z.read(self.rels(part)[rid]))).convert('RGBA'))


def setup_font():
    for file in ['C:/Windows/Fonts/msyh.ttc', 'C:/Windows/Fonts/simhei.ttf']:
        if Path(file).exists():
            font_manager.fontManager.addfont(file)
            name = font_manager.FontProperties(fname=file).get_name()
            plt.rcParams['font.family'] = name
            break
    else:
        plt.rcParams['font.family'] = ['Noto Sans CJK SC', 'SimHei', 'DejaVu Sans']
    plt.rcParams.update({'font.size': 9, 'text.color': WHITE,
                         'axes.labelcolor': WHITE, 'xtick.color': WHITE,
                         'ytick.color': WHITE, 'axes.unicode_minus': False,
                         'svg.fonttype': 'path'})


def native_size(shapes):
    """由图内形状的 EMU 尺寸和相对锚点反算 Excel 图表英寸尺寸。"""
    if shapes is not None:
        for a in shapes:
            ext = a.find('./sp/spPr/xfrm/ext')
            if ext is not None and a.find('./from/x') is not None:
                dx = float(a.findtext('./to/x')) - float(a.findtext('./from/x'))
                dy = float(a.findtext('./to/y')) - float(a.findtext('./from/y'))
                if dx > 0 and dy > 0:
                    return int(ext.get('cx')) / dx / 914400, int(ext.get('cy')) / dy / 914400
    return 5.667, 4.25


def text_shape(fig, shape, box):
    """按原文段落、字号、颜色、对齐及行距绘制文本框。"""
    x, y, w, h = box  # y 从顶部起算
    body = shape.find('txBody')
    if body is None:
        return
    paras = body.findall('p')
    bp = body.find('bodyPr')
    pad_x = float(bp.get('lIns', 91440)) / 914400 / fig.get_figwidth() if bp is not None else 0
    pad_y = float(bp.get('tIns', 45720)) / 914400 / fig.get_figheight() if bp is not None else 0
    cursor = 1 - y - pad_y
    renderer = fig.canvas.get_renderer()
    for p in paras:
        runs = p.findall('r') + p.findall('fld')
        text = ''.join(r.findtext('t', '') for r in runs)
        if not text:
            continue
        rp = runs[0].find('rPr')
        fs = float(rp.get('sz', 1100)) / 100 if rp is not None else 11
        bold = rp is not None and rp.get('b') == '1'
        color = fill(rp)
        align = attr(p, 'pPr', 'l', key='algn')
        ha = {'l': 'left', 'ctr': 'center', 'r': 'right'}.get(align, 'left')
        tx = x + pad_x if ha == 'left' else x + (w / 2 if ha == 'center' else w - pad_x)
        # Excel 会在文本框边缘换行；按字体真实宽度逐字进行同样处理。
        prop = font_manager.FontProperties(family=plt.rcParams['font.family'], size=fs,
                                           weight='bold' if bold else 'normal')
        max_pixels = (w - 2 * pad_x) * fig.bbox.width
        lines, line = [], ''
        for ch in text:
            test = line + ch
            width = renderer.get_text_width_height_descent(test, prop, False)[0]
            if width > max_pixels and line:
                lines.append(line)
                line = ch
            else:
                line = test
        lines.append(line)
        spacing = value(p, './pPr/lnSpc/spcPts', fs * 120) / 100
        anchor = bp.get('anchor', 't') if bp is not None else 't'
        if anchor == 'ctr' and len(paras) == 1:
            cursor = 1 - y - h / 2
        for line in lines:
            run_colors = [fill(r.find('rPr')) for r in runs]
            va = 'center' if anchor == 'ctr' and len(paras) == 1 else 'top'
            if len(set(run_colors)) > 1 and len(lines) == 1:
                widths = [renderer.get_text_width_height_descent(r.findtext('t',''),prop,False)[0]
                          / fig.bbox.width for r in runs]
                px = tx - (sum(widths)/2 if ha == 'center' else sum(widths) if ha == 'right' else 0)
                for r,rc,rw in zip(runs,run_colors,widths):
                    fig.text(px,cursor,r.findtext('t',''),fontsize=fs,color=rc,ha='left',va=va,
                             weight='bold' if bold else 'normal',zorder=15)
                    px += rw
            else:
                fig.text(tx, cursor, line, fontsize=fs, color=color, ha=ha, va=va,
                         weight='bold' if bold else 'normal', zorder=15)
            cursor -= spacing / 72 / fig.get_figheight()


def draw_shapes(fig, root, skip_empty=False):
    if root is None:
        return
    for a in root:
        shape = a.find('sp')
        if shape is None or a.find('./from/x') is None:
            continue
        x, y = float(a.findtext('./from/x')), float(a.findtext('./from/y'))
        w, h = float(a.findtext('./to/x')) - x, float(a.findtext('./to/y')) - y
        sp = shape.find('spPr')
        fc, ec = fill(sp, '#00000000'), fill(sp.find('ln'), '#00000000')
        geom = attr(sp, 'prstGeom', 'rect', key='prst')
        if not skip_empty and (fc[3] or ec[3]):
            if geom == 'roundRect':
                patch = FancyBboxPatch((x, 1-y-h), w, h, boxstyle='round,pad=0,rounding_size=0.008',
                                       facecolor=fc, edgecolor=ec, linewidth=.75, transform=fig.transFigure, zorder=4)
            elif geom == 'rect':
                patch = Rectangle((x, 1-y-h), w, h, facecolor=fc, edgecolor=ec,
                                  linewidth=.75, transform=fig.transFigure, zorder=4)
            else:
                patch = None
            if patch is not None:
                fig.add_artist(patch)
        text_shape(fig, shape, (x, y, w, h))


def axes_for(fig, root):
    layout = root.find('./chart/plotArea/layout/manualLayout')
    x, y, w, h = [value(layout, k, d) for k, d in zip('xywh', [.12, .28, .8, .56])]
    ax = fig.add_axes([x, 1-y-h, w, h], facecolor=BG)
    for spine in ax.spines.values():
        spine.set_visible(False)
    ax.tick_params(length=0, pad=5)
    ax.set_axisbelow(True)
    return ax


def vertical_axes(ax, root, labels, maximum, step=None):
    ax.set_xlim(-.5, len(labels)-.5)
    ax.set_ylim(0, maximum)
    ax.set_xticks(range(len(labels)))
    ax.set_xticklabels(labels)
    va = root.find('./chart/plotArea/valAx')
    if attr(va, 'delete', '0') == '1':
        ax.set_yticks([])
    else:
        if step:
            ax.set_yticks(np.arange(0, maximum + step * .01, step))
        if va.find('majorGridlines') is not None:
            ax.grid(axis='y', color=WHITE, alpha=.2, linestyle=(0, (6, 4)), linewidth=.5)
    ca = root.find('./chart/plotArea/catAx')
    if ca is not None and ca.find('./spPr/ln/solidFill') is not None:
        ax.spines['bottom'].set_visible(True)
        ax.spines['bottom'].set_color(fill(ca.find('./spPr/ln')))
        ax.spines['bottom'].set_linewidth(.5)


def labels_above(ax, x, values, offset=4, fontsize=9, color=WHITE):
    for xx, yy in zip(x, values):
        ax.annotate(f'{yy:g}', (xx, yy), xytext=(0, offset), textcoords='offset points',
                    ha='center', va='bottom', fontsize=fontsize, color=color)


def gradient(ax, extent, top, bottom, clip=None):
    rgba = np.linspace(colors.to_rgba(bottom), colors.to_rgba(top), 256).reshape(256, 1, 4)
    im = ax.imshow(rgba, extent=extent, origin='lower', aspect='auto', interpolation='bicubic', zorder=2)
    if clip is not None:
        im.set_clip_path(clip)


def smooth_curve(y):
    """无 SciPy 依赖的三次 Hermite 插值，通过全部原始点。"""
    y = np.asarray(y, float)
    slopes = np.gradient(y)
    xx, yy = [], []
    for i in range(len(y)-1):
        t = np.linspace(0, 1, 50)
        v = (2*t**3-3*t**2+1)*y[i] + (t**3-2*t**2+t)*slopes[i]
        v += (-2*t**3+3*t**2)*y[i+1] + (t**3-t**2)*slopes[i+1]
        xx.extend(i+t)
        yy.extend(v)
    return xx, yy


def render(source, number, dpi):
    root = source.chart(number) if number != 7 else None
    shapes = source.shapes(number) if number != 7 else None
    size = native_size(shapes) if number not in (7, 15) else (5.667, 4.25)
    fig = plt.figure(figsize=size, dpi=dpi, facecolor=BG)
    s = source.book.worksheets[number-1]
    if number in (7, 15):
        ax = fig.add_axes([.10, .18, .80, .52], facecolor=BG)
        ax.set_axis_off()
    else:
        ax = axes_for(fig, root)
    cat = source.col(number, 2) if number not in (10, 14, 15) else []
    vals = source.col(number, 3) if number not in (10, 14, 15) else []
    x = np.arange(len(cat))

    if number in (1, 2, 3, 4):
        maximum = {1:4000, 2:4000, 3:1200, 4:10000}[number]
        step = {1:1000, 2:1000, 3:200, 4:2000}[number]
        vertical_axes(ax, root, cat, maximum, step)
        width = 1/(1+value(root, './/barChart/gapWidth', 150)/100)
        if number == 1:
            for xx, yy in zip(x, vals):
                gradient(ax, (xx-width/2, xx+width/2, 0, yy), BLUE, '#00B0F0')
            labels_above(ax, x, vals)
        elif number == 2:
            ax.bar(x, vals, width, color=BLUE)
            labels_above(ax, x[:-1], vals[:-1])
            labels_above(ax, x[-1:], vals[-1:], offset=-12)
            mean = float(np.mean(vals))
            ax.plot(x, np.repeat(mean, len(x)), color='#FFC000', linewidth=1.5)
        elif number == 3:
            points = root.findall('.//barChart/ser/dPt')
            for xx, yy, point in zip(x, vals, points):
                rid = attr(point, './spPr/blipFill/blip', key='embed')
                ax.imshow(source.picture(3, rid), extent=(xx-width/2, xx+width/2, 0, yy),
                          aspect='auto', interpolation='bicubic', zorder=2)
            # 原图的标签被手动移到各自半透明圆角框内。
            boxes = sorted([a for a in shapes if attr(a,'./sp/spPr/prstGeom',key='prst')=='roundRect'],
                           key=lambda a:float(a.findtext('./from/x')))
            for a,v in zip(boxes,vals):
                xx=(float(a.findtext('./from/x'))+float(a.findtext('./to/x')))/2
                yy=1-(float(a.findtext('./from/y'))+float(a.findtext('./to/y')))/2
                fig.text(xx,yy,str(v),fontsize=9,ha='center',va='center',zorder=15)
        else:
            palette = ['#7BBDD5','#49A098','#E66B4C','#FFC000','#0097E0',BLUE,'#4A5BD1','#464CAC']
            ax.bar(x, vals, width, color=palette)
            for xx, yy, c in zip(x, vals, palette):
                ax.annotate(str(yy), (xx, yy), xytext=(0, 11), textcoords='offset points',
                            ha='center', fontsize=8, color=WHITE,
                            bbox=dict(boxstyle='square,pad=.22', fc=c, ec='none'),
                            arrowprops=dict(arrowstyle='wedge,tail_width=.4', fc=c, ec=c,
                                            shrinkA=0, shrinkB=0, relpos=(.5,0)))
    elif number == 5:
        profit = source.col(5,4)
        vertical_axes(ax, root, cat, 6000)
        width = 1/(2-.3+2.19)
        delta = width*.7/2
        ax.bar(x-delta, vals, width, color=BLUE)
        ax.bar(x+delta, profit, width, color=PINK)
        labels_above(ax, x-delta, vals)
        labels_above(ax, x+delta, profit)
    elif number == 6:
        left, right = source.col(6,4), source.col(6,6)
        # 原图通过两个占位序列将左侧条形右对齐，中间保留区域名称。
        base = np.array(source.col(6,3))
        gap = np.array(source.col(6,5))
        ax.barh(x, left, left=base, height=.4, color=BLUE)
        ax.barh(x, right, left=base+left+gap, height=.4, color=PINK)
        for i in x:
            ax.text(base[i]+45, i, str(left[i]), ha='left', va='center')
            ax.text(base[i]+left[i]+gap[i]/2, i, cat[i], ha='center', va='center')
            ax.text(base[i]+left[i]+gap[i]+right[i]-45, i, str(right[i]), ha='right', va='center')
        ax.set_xlim(0,7000)
        ax.set_ylim(-.5,4.5)
        ax.set_xticks([]); ax.set_yticks([])
    elif number == 7:
        # Excel 第 7 页使用 REPT("|", 数值*200) 模拟条形，无 Chart 对象。
        right = source.col(7,4)
        gap = .055
        ax.set_xlim(-.53,.53); ax.set_ylim(4.7,-.8)
        for i,(name,l,r) in enumerate(zip(cat, vals, right)):
            ax.barh(i, l, left=-gap-l, height=.5, color=BLUE)
            ax.barh(i, r, left=gap, height=.5, color=PINK)
            ax.text(0,i,name,ha='center',va='center',fontsize=11)
            ax.text(-gap-l-.012,i,f'{l:.0%}',ha='right',va='center',fontsize=10)
            ax.text(gap+r+.012,i,f'{r:.0%}',ha='left',va='center',fontsize=10)
        drawing = source.root('xl/drawings/drawing13.xml')
        textshapes = [e for e in drawing.findall('./twoCellAnchor/sp')
                      if any('目标完成情况' in (t.text or '') for t in e.findall('.//t'))]
        if textshapes:
            text_shape(fig,textshapes[0],(.045,.055,.92,.22))
        notes = [e for e in drawing.findall('./twoCellAnchor/sp') if any('注：' in (t.text or '') for t in e.findall('.//t'))]
        if notes:
            text_shape(fig,notes[0],(.035,.92,.94,.065))
    elif number == 8:
        remaining = np.array(source.col(8,4))
        gap = source.col(8,5)
        yoy = source.col(8,6)
        ax.barh(x,vals,height=1/1.3,color='#09387E')
        ax.barh(x,remaining,left=vals,height=1/1.3,color='#82ADD7')
        ax.barh(x,gap,left=np.array(vals)+remaining,height=1/1.3,color='#9B3D4F')
        for i,v in enumerate(vals):
            ax.text(v-60,i,str(v),ha='right',va='center',fontsize=10)
            ax.text(v+remaining[i]+gap[i]/2,i,f'{yoy[i]:.1%}',ha='center',va='center',fontsize=10)
        ax.set_yticks(x); ax.set_yticklabels(cat)
        ax.set_ylim(-.5,5.5); ax.set_xlim(0,6000); ax.set_xticks([])
    elif number == 9:
        current = source.col(9,4)
        vertical_axes(ax,root,cat,6000)
        width=1/(2+.27+2.19); delta=width*1.27/2
        ax.bar(x-delta,vals,width,color=BLUE)
        ax.bar(x+delta,current,width,color='#82ADD7')
        labels_above(ax,x-delta,vals)
        labels_above(ax,x+delta,current,offset=-13)
        for i,(old,new) in enumerate(zip(vals,current)):
            ax.annotate('',(i+delta,new),(i+delta,old),
                        arrowprops=dict(arrowstyle='-|>',lw=.75,color=WHITE,shrinkA=0,shrinkB=0))
            ax.plot([i-delta,i+delta],[old,old],color=BLUE,lw=.75)
            ax.text(i+delta+.075,old+80,str(old-new),ha='left',fontsize=9)
        fig.legend(handles=[Rectangle((0,0),1,1,color=BLUE,label='2021销量'),
                            Rectangle((0,0),1,1,color='#82ADD7',label='2022销量')],
                   loc='upper left',bbox_to_anchor=(.05,.74),frameon=False,ncol=2,fontsize=9,
                   borderaxespad=0,handlelength=1,handletextpad=.4,columnspacing=1)
    elif number == 10:
        cat=source.col(10,2,4,10)
        starts=source.col(10,3,4,10); days=source.col(10,4,4,10); complete=source.col(10,5,4,10)
        # 原图只有“项目天数”系列；完成百分比是标签，并无第二条进度条。
        epoch=datetime(1899,12,30)
        start=np.array([(d-epoch).days for d in starts])
        ax.barh(range(7),days,left=start,height=.5,color=BLUE)
        for i,(v,d,p) in enumerate(zip(start,days,complete)):
            ax.text(v+d/2,i,f'{p:.0%}',ha='center',va='center',fontsize=9)
        ax.set_yticks(range(7));ax.set_yticklabels(cat)
        ax.set_ylim(6.5,-.5);ax.set_xlim(44621,44726)
        ticks=np.arange(44621,44727,15)
        ax.set_xticks(ticks)
        ax.set_xticklabels([(epoch+timedelta(days=int(v))).strftime('%m/%d/%Y').lstrip('0') for v in ticks],fontsize=7)
        ax.xaxis.tick_top()
        ax.grid(axis='x',color=WHITE,alpha=.4,lw=.5,ls=(0,(6,4)))
    elif number == 11:
        months=source.col(11,3); data=source.col(11,4)
        vertical_axes(ax,root,months,4000,1000)
        xx,yy=smooth_curve(data)
        ax.plot(xx,yy,color='#E66B4C',lw=1)
        peak=int(np.argmax(data))
        labels_above(ax,[peak],[data[peak]])
    elif number == 12:
        vertical_axes(ax,root,cat,.8)
        ax.vlines(x,0,vals,color=PINK,linewidth=.75)
        ax.plot(x,vals,linestyle='none',marker='D',ms=8,mfc=BG,mec=PINK,mew=.75)
        for xx,yy in zip(x,vals):
            ax.annotate(f'{yy:.2%}',(xx,yy),xytext=(0,8),textcoords='offset points',ha='center',fontsize=9)
    elif number == 13:
        newer=source.col(13,4)
        vertical_axes(ax,root,cat,3000,500)
        for data,c in [(vals,PINK),(newer,BLUE)]:
            ax.plot(x,data,color=c,lw=1.5,marker='o',ms=5,mec=BG,mew=.75)
        # 对比两条线的位置，沿原图上下错开标签以避免碰撞。
        for i,(a,b) in enumerate(zip(vals,newer)):
            for v,other in [(a,b),(b,a)]:
                ax.annotate(str(v),(i,v),xytext=(0,7 if v>other else -13),
                            textcoords='offset points',ha='center',fontsize=9)
        fig.legend(handles=[Line2D([],[],color=PINK,marker='o',ms=4,label='2021年'),
                            Line2D([],[],color=BLUE,marker='o',ms=4,label='2022年')],
                   loc='upper left',bbox_to_anchor=(.64,.73),frameon=False,ncol=2,fontsize=9,
                   borderaxespad=0,handlelength=1.1,columnspacing=1)
    elif number == 14:
        rate=float(s['B3'].value)
        if not 0 <= rate <= 1:
            raise ValueError('圆环完成率必须在 0~1 之间')
        ax.set_axis_off();ax.set_xlim(-1,1);ax.set_ylim(-1,1);ax.set_aspect('equal')
        # Excel 首扇区角度从 12 点顺时针计，300 度对应 Matplotlib 的 150 度。
        start=90-value(root,'.//firstSliceAng',300)
        ring=Wedge((0,0),1,start-360*rate,start,width=.1,facecolor='none',edgecolor='none')
        ax.add_patch(Wedge((0,0),1,0,360,width=.1,facecolor=(.95,.95,.95,.1),edgecolor='none'))
        ax.add_patch(ring)
        gradient(ax,(-1,1,-1,1),'#7030A0',PINK,ring)
        ax.set_aspect('equal')
    elif number == 15:
        rate=float(s['B3'].value)
        if not 0 <= rate <= 1:
            raise ValueError('水球完成率必须在 0~1 之间')
        # 原图为圆形图片叠加裁切，液面是水平线，不额外添加波浪。
        ax.set_position([.30,.17,.42,.56]);ax.set_xlim(-1.05,1.05);ax.set_ylim(-1.05,1.05);ax.set_aspect('equal')
        circle=Circle((0,0),1,facecolor='none',edgecolor=BLUE,lw=1.5)
        ax.add_patch(circle)
        water=Rectangle((-1,-1),2,2*rate,facecolor=BLUE,edgecolor='none')
        ax.add_patch(water);water.set_clip_path(circle)
        ax.text(0,0,f'{rate:.0%}',fontsize=40,ha='center',va='center')
        drawing=source.root('xl/drawings/drawing28.xml')
        for shape in drawing.findall('./twoCellAnchor/sp'):
            texts=''.join(t.text or '' for t in shape.findall('.//t'))
            if '年上半年目标完成率' in texts:
                text_shape(fig,shape,(.045,.045,.91,.23))
            elif '注：' in texts:
                text_shape(fig,shape,(.035,.91,.85,.07))

    if number not in (7,15):
        # 圆角图标签周围的描边也是原图形状的一部分。
        draw_shapes(fig, shapes)
    return fig


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input',type=Path)
    parser.add_argument('--output',type=Path,default=HERE/'生成图表')
    parser.add_argument('--dpi',type=int,default=240)
    args=parser.parse_args()
    candidates=[HERE/FILENAME,HERE.parent/FILENAME,Path.cwd()/FILENAME,
                Path(r'D:/CP/python可视化课程/《Excel数据可视化 - 从图表到数据大屏》-清华-郭宏远')/FILENAME]
    path=args.input or next((p for p in candidates if p.exists()),None)
    if path is None or not path.exists():
        parser.error('找不到工作簿，请用 --input 指定完整路径。')
    setup_font()
    source=Source(path)
    if len(source.book.worksheets)<15:
        raise ValueError('输入文件需要包含前 15 个示例工作表。')
    args.output.mkdir(parents=True,exist_ok=True)
    files=[]
    for i,s in enumerate(source.book.worksheets[:15],1):
        fig=render(source,i,args.dpi)
        name=f'{i:02d}_{s.title.split(" ",1)[-1]}'
        target=args.output/(name+'.png')
        fig.savefig(target,dpi=args.dpi,facecolor=BG)
        fig.savefig(args.output/(name+'.svg'),facecolor=BG)
        plt.close(fig)
        files.append(target)
        print(f'已生成 {i:02d}/15：{target.name}')
    # 总览图按 3 列 × 5 行排列，单张高清图另存，保持原始长宽比。
    overview=Image.new('RGB',(1800,5*490),'#10142D')
    draw=ImageDraw.Draw(overview)
    font_path=Path('C:/Windows/Fonts/msyh.ttc')
    font=ImageFont.truetype(str(font_path),18) if font_path.exists() else ImageFont.load_default()
    for i,file in enumerate(files):
        im=Image.open(file).convert('RGB')
        im.thumbnail((580,445),Image.Resampling.LANCZOS if hasattr(Image,'Resampling') else Image.LANCZOS)
        x=(i%3)*600+(600-im.width)//2;y=(i//3)*490+30
        overview.paste(im,(x,y))
        draw.text(((i%3)*600+12,(i//3)*490+6),file.stem,font=font,fill='white')
    overview.save(args.output/'00_全部15图总览.png')
    manifest={'source':str(path.resolve()),'charts':len(files),'renderer':'Matplotlib',
              'outputs':[p.name for p in files],
              'note':'由单元格数据与原工作簿图表样式重绘；不是 Excel 原生导出。'}
    (args.output/'生成记录.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    source.book.close();source.z.close()
    print('完成。输出目录：'+str(args.output.resolve()))


if __name__=='__main__':
    main()
